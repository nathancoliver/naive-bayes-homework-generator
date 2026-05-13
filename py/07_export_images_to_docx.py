# Import docx NOT python-docx
import docx
  
# Create an instance of a word document
doc = docx.Document('C:/Users/nathan.oliver/Desktop/Python/IG_Physics_Question_Bank/docx/test/test1.docx')
  
# # Add a Title to the document
# doc.add_heading('GeeksForGeeks', 0)
  
# # Image in its native size
# doc.add_heading('Image in Native Size:', 3)
doc.add_picture('C:/Users/nathan.oliver/Desktop/Python/IG_Physics_Question_Bank/pdf/test/test1.png',width=5000000)
doc.add_picture('C:/Users/nathan.oliver/Desktop/Python/IG_Physics_Question_Bank/pdf/test/test1.png')

  
# Now save the document to a location
doc.save('C:/Users/nathan.oliver/Desktop/Python/IG_Physics_Question_Bank/docx/test/test1.docx')
